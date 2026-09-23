"""Build editable Word and printable PDF handoff from Markdown sources."""
from pathlib import Path
from html import escape
import argparse
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as RasterImage
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, KeepTogether, Table, TableStyle, Spacer, Image

ROOT = Path(__file__).resolve().parents[1]
SOURCES = [ROOT/'docs'/name for name in ['HANDOFF.md', 'STATUT-LIVRAISON.md', 'CONTRACT.md', 'REFERENTIEL.md', 'MODELE-DONNEES.md', 'NOTICE.md', 'CONNEXION-JUGES.md', 'MATRICE-DROITS.md', 'EXPLOITATION.md', 'RECETTE.md', 'PV-RECETTE.md', 'DESIGN.md', 'LANCEMENT.md']] + [ROOT/'README-OPERATIONS.md'] + sorted((ROOT/'docs').glob('REPORT-*.md'))


def clean(text):
    return text.replace('`', '').replace('**', '').replace('\n', ' ').strip()


def cells(line):
    return [clean(cell.replace(r'\|', '|')) for cell in re.split(r'(?<!\\)\|', line.strip().strip('|'))]


def blocks(source):
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith('```'):
            index += 1
            continue
        if line.startswith('|') and index+1 < len(lines) and re.fullmatch(r'[\s|:\-]+', lines[index+1]):
            rows = [cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = cells(lines[index])
                if len(row) != len(rows[0]):
                    raise ValueError('Nombre de colonnes incohérent : '+lines[index])
                rows.append(row)
                index += 1
            yield 'table', rows
            continue
        picture = re.fullmatch(r'!\[([^]]*)\]\((.+?)\)(?:\s*)', line)
        if picture:
            reference = re.sub(r'\s+"[^"]*"$', '', picture[2]).strip().strip('<>')
            yield 'image', (clean(picture[1]), reference)
            index += 1
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)', line)
        if heading:
            yield 'title' if len(heading[1]) == 1 else 'heading', clean(heading[2])
            index += 1
            continue
        paragraph = [line]
        index += 1
        if not re.match(r'^([-*] |\d+\. )', line):
            while index < len(lines) and lines[index].strip() and not re.match(r'^(#|!\[|\||[-*] |\d+\. |```)', lines[index].strip()):
                paragraph.append(lines[index].strip())
                index += 1
        yield 'text', clean(' '.join(paragraph))


def unicode_font():
    candidates = [os.environ.get('FIBDA_DOC_FONT', ''), '/Library/Fonts/Arial Unicode.ttf', '/System/Library/Fonts/Supplemental/Arial Unicode.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 'C:/Windows/Fonts/arial.ttf']
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise RuntimeError('Police Unicode absente : définir FIBDA_DOC_FONT vers un fichier TTF Unicode.')


def build(sources, output):
    output.mkdir(parents=True, exist_ok=True)
    font_path = unicode_font()
    pdfmetrics.registerFont(TTFont('FIBDAUnicode', font_path))
    pdfmetrics.registerFontFamily('FIBDAUnicode', normal='FIBDAUnicode', bold='FIBDAUnicode', italic='FIBDAUnicode', boldItalic='FIBDAUnicode')
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.1)
    for key in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
        word.styles[key].font.name = 'Arial'
        word.styles[key].font.color.rgb = RGBColor(0, 0, 0)
    word.styles['Normal'].font.size = Pt(10)
    word.styles['Normal'].paragraph_format.space_after = Pt(7)
    word.styles['Normal'].paragraph_format.keep_together = True
    for element in word.styles.element.iter():
        for child in list(element):
            if child.tag == qn('w:pBdr'):
                element.remove(child)
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = 'FIBDAUnicode'
    styles['Heading2'].keepWithNext = True
    styles['Title'].keepWithNext = True
    styles['Normal'].fontSize, styles['Normal'].leading, styles['Normal'].spaceAfter = 10, 14, 8
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=8.5, leading=11, spaceAfter=0, splitLongWords=True)
    caption_style = ParagraphStyle('Caption', parent=styles['Normal'], fontSize=8.5, leading=11, alignment=1, spaceBefore=5, spaceAfter=12)
    header_style = ParagraphStyle('Header', parent=cell_style, textColor=colors.white)
    story = []
    width = A4[0]-84
    for index, filename in enumerate(sources):
        if index:
            story.append(PageBreak())
        pending = []
        for kind, value in blocks(filename.read_text()):
            if kind in ('title', 'heading'):
                heading = word.add_paragraph(value, 'Title' if kind == 'title' else 'Heading 1')
                if kind == 'title' and index:
                    heading.paragraph_format.page_break_before = True
                pending.append(Paragraph(escape(value), styles['Title' if kind == 'title' else 'Heading2']))
            elif kind == 'image':
                caption, reference = value
                if '://' in reference or Path(reference).is_absolute():
                    raise ValueError('Une capture doit être un fichier relatif au document : '+reference)
                picture_path = (filename.parent/reference).resolve()
                with RasterImage.open(picture_path) as picture:
                    pixel_width, pixel_height = picture.size
                scale = min(width/pixel_width, 430/pixel_height)
                draw_width, draw_height = pixel_width*scale, pixel_height*scale
                flowables = pending + [Image(str(picture_path), width=draw_width, height=draw_height)]
                if caption:
                    flowables.append(Paragraph(escape(caption), caption_style))
                story.append(KeepTogether(flowables))
                pending = []
                image_paragraph = word.add_paragraph()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.paragraph_format.keep_with_next = bool(caption)
                # Word has a slightly narrower text area; preserve the same aspect ratio.
                word_scale = min(Cm(16.8).pt/draw_width, 1)
                image_paragraph.add_run().add_picture(str(picture_path), width=Pt(draw_width*word_scale), height=Pt(draw_height*word_scale))
                if caption:
                    legend = word.add_paragraph(caption, 'Caption')
                    legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif kind == 'table':
                count = len(value[0])
                # Two-column registers give descriptions the width they need; forms retain a signature column.
                fractions = [0.26, 0.74] if count == 2 else [0.43, 0.37, 0.20] if count == 3 else [1/count]*count
                table = word.add_table(rows=0, cols=count)
                table.style = 'Table Grid'
                table.autofit = False
                for column, fraction in zip(table.columns, fractions):
                    column.width = Cm(16.8*fraction)
                for row_index, row in enumerate(value):
                    cells_word = table.add_row().cells
                    for cell, text, fraction in zip(cells_word, row, fractions):
                        cell.width = Cm(16.8*fraction)
                        cell.text = text
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(4)
                            paragraph.paragraph_format.keep_with_next = row_index == 0
                            for run in paragraph.runs:
                                run.font.size = Pt(8.5)
                                run.bold = row_index == 0
                    properties = table.rows[-1]._tr.get_or_add_trPr()
                    properties.append(OxmlElement('w:cantSplit'))
                    if row_index == 0:
                        properties.append(OxmlElement('w:tblHeader'))
                pdf_table = Table([[Paragraph(escape(cell), header_style if row_index == 0 else cell_style) for cell in row] for row_index, row in enumerate(value)], colWidths=[width*f for f in fractions], repeatRows=1, hAlign='LEFT')
                pdf_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#164b3d')),('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#b8c5bf')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
                story.extend(pending)
                pending = []
                story.extend([pdf_table, Spacer(1, 10)])
                word.add_paragraph()
            else:
                word.add_paragraph(value)
                story.append(KeepTogether(pending + [Paragraph(escape(value), styles['Normal'])]))
                pending = []
        story.extend(pending)
    word.core_properties.title = 'Dossier développeur et exploitation FIBDA'
    word.core_properties.author = 'FIBDA'
    word.save(output/'Dossier-developpeur-FIBDA.docx')
    def footer(canvas, doc):
        canvas.setFont('FIBDAUnicode', 8)
        canvas.drawString(42,25,'FIBDA — Dossier développeur et exploitation')
        canvas.drawRightString(A4[0]-42,25,str(doc.page))
    SimpleDocTemplate(str(output/'Dossier-developpeur-FIBDA.pdf'), pagesize=A4,leftMargin=42,rightMargin=42,topMargin=42,bottomMargin=42,title='Dossier développeur et exploitation FIBDA').build(story,onFirstPage=footer,onLaterPages=footer)
    print(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', type=Path, default=ROOT/'livrables')
    parser.add_argument('--sources', nargs='+', type=Path, default=SOURCES)
    arguments = parser.parse_args()
    build(arguments.sources, arguments.output_dir)
