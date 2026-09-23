"""Bounded, non-executing imports and image normalization."""
import csv
import io
import warnings
import zipfile
from pathlib import PurePosixPath

MAX_FILE = 20 * 1024 * 1024
MAX_TOTAL = 200 * 1024 * 1024
MAX_ROWS = 10000
FIELDS = {'first_name','last_name','birth_date','sex','nationalities','country','club','section','height_cm','weight_kg','category_id'}


def safe_members(archive, max_total=MAX_TOTAL):
    infos = archive.infolist()
    if len(infos) > 2000 or sum(x.file_size for x in infos) > max_total:
        raise ValueError('Archive trop volumineuse')
    seen = set()
    for info in infos:
        path = PurePosixPath(info.filename)
        if (path.is_absolute() or '..' in path.parts or '\\' in info.filename or ':' in info.filename
                or info.filename in seen or ((info.external_attr >> 16) & 0o170000) == 0o120000):
            raise ValueError('Chemin dangereux ou doublon dans archive')
        seen.add(info.filename)
        if info.file_size > MAX_FILE:
            raise ValueError('Fichier trop volumineux')
    return infos


def preview_import(data, filename):
    if len(data) > MAX_FILE:
        raise ValueError('Fichier trop volumineux')
    errors, rows, notices = [], [], []
    if filename.lower().endswith('.csv'):
        text = data.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        raw = list(csv.reader(io.StringIO(text), dialect))
    elif filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            safe_members(z)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        raw = []
        for index, row in enumerate(workbook.active.iter_rows(values_only=True)):
            if index > MAX_ROWS:
                raise ValueError('Maximum 10000 lignes')
            from datetime import date, datetime
            raw.append([value.date().isoformat() if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else value for value in row])
        workbook.close()
    else:
        raise ValueError('Formats acceptés : CSV UTF-8 et XLSX')
    if not raw:
        return {'rows': [], 'errors': [{'row': 1, 'message': 'Fichier vide'}], 'warnings': []}
    headers = [str(x or '').strip() for x in raw[0]]
    required = {'first_name','last_name','birth_date','sex','section','country','nationalities'}
    if len(headers) != len(set(headers)) or not required.issubset(headers):
        errors.append({'row': 1, 'message': 'Entêtes uniques obligatoires : first_name, last_name, birth_date, sex, section, country, nationalities'})
    unknown = set(headers) - FIELDS
    if unknown:
        errors.append({'row': 1, 'message': 'Colonnes inconnues : ' + ', '.join(sorted(unknown))})
    if len(raw) > MAX_ROWS + 1:
        raise ValueError('Maximum 10000 lignes')
    seen = set()
    for n, values in enumerate(raw[1:], 2):
        if not any(x is not None and str(x).strip() for x in values):
            continue
        row = {k: str(v or '').strip() for k, v in zip(headers, values)}
        if len(values) != len(headers):
            errors.append({'row': n, 'message': 'Nombre de colonnes incohérent'})
        if any(str(v).lstrip().startswith(('=', '+', '-', '@')) for v in values if v is not None):
            errors.append({'row': n, 'message': 'Formule ou préfixe exécutable refusé'})
        if not row.get('first_name') or not row.get('last_name'):
            errors.append({'row': n, 'message': 'Nom et prénom obligatoires'})
        if row.get('sex') not in ('M', 'F'):
            errors.append({'row': n, 'message': 'Sexe attendu M ou F'})
        if row.get('section') not in ('amateur', 'pro'):
            errors.append({'row': n, 'message': 'Section attendue amateur ou pro'})
        for key in ('height_cm', 'weight_kg'):
            if row.get(key):
                try:
                    from decimal import Decimal, InvalidOperation
                    value = Decimal(row[key].replace(',', '.'))
                    if not value.is_finite() or value <= 0:
                        raise InvalidOperation
                    row[key] = str(value)
                except InvalidOperation:
                    errors.append({'row': n, 'message': key + ' doit être un nombre positif'})
        if not row.get('birth_date'):
            errors.append({'row': n, 'message': 'Date de naissance obligatoire'})
        if row.get('birth_date'):
            try:
                from datetime import date
                date.fromisoformat(row['birth_date'])
            except ValueError:
                errors.append({'row': n, 'message': 'Date attendue AAAA-MM-JJ'})
        key = (row.get('first_name', '').casefold(), row.get('last_name', '').casefold(), row.get('birth_date'))
        if key in seen:
            notices.append({'row': n, 'message': 'Identité répétée : vérifier avant import'})
        seen.add(key)
        if row.get('nationalities'):
            row['nationalities'] = [x.strip().upper() for x in row['nationalities'].split('|') if x.strip()]
        countries = [row.get('country', '')] + row.get('nationalities', []) if isinstance(row.get('nationalities'), list) else []
        if len(countries) < 2 or any(len(c) != 2 or not c.isascii() or not c.isalpha() or c != c.upper() for c in countries):
            errors.append({'row': n, 'message': 'Pays et nationalités obligatoires : codes de deux lettres majuscules ; séparer les nationalités par |'})
        rows.append(row)
    return {'rows': rows, 'errors': errors, 'warnings': notices}


def sanitize_photo(data, crop=None):
    from PIL import Image, ImageOps
    if len(data) > MAX_FILE:
        raise ValueError('Photo trop volumineuse')
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        try:
            with Image.open(io.BytesIO(data)) as source:
                if source.format not in ('JPEG', 'PNG', 'WEBP') or source.width * source.height > 25000000:
                    raise ValueError('Photo non supportée ou dimensions excessives')
                source.load()
                picture = ImageOps.exif_transpose(source).convert('RGB')
        except (OSError, Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
            raise ValueError('Photo invalide') from exc
    if crop is not None:
        if len(crop) != 4 or not all(isinstance(x, (int, float)) for x in crop):
            raise ValueError('Recadrage attendu [gauche,haut,droite,bas] en pixels')
        left, top, right, bottom = crop
        if not (0 <= left < right <= picture.width and 0 <= top < bottom <= picture.height):
            raise ValueError('Recadrage hors photo')
        picture = picture.crop(crop)
    picture.thumbnail((1600, 2000))
    output = io.BytesIO()
    picture.save(output, 'JPEG', quality=88, optimize=True)
    return output.getvalue()


def batch_photos(data, mappings):
    if len(data) > MAX_TOTAL:
        raise ValueError('Archive trop volumineuse')
    result, owners = [], set()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        available = {x.filename for x in safe_members(archive) if not x.is_dir()}
        for mapping in mappings:
            name = mapping.get('filename')
            owner = (mapping.get('owner_type'), mapping.get('owner_id'), mapping.get('kind'))
            if name not in available or owner[0] not in ('person', 'official') or not owner[1] or owner[2] not in ('portrait', 'full') or owner in owners:
                raise ValueError('Association photo absente, invalide ou répétée')
            owners.add(owner)
            result.append(dict(owner_type=owner[0], owner_id=owner[1], kind=owner[2], data=sanitize_photo(archive.read(name), mapping.get('crop'))))
    return result
